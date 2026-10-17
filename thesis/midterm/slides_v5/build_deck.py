"""Rebuild the midterm deck around the v5 framework (English slides and notes).

Usage: python build_deck.py <orig.pptx> <img_dir> <out.pptx>
"""
import copy
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC, IMG, OUT = sys.argv[1:4]

RED = RGBColor(0x92, 0x0C, 0x13)
ROSE = RGBColor(0xE8, 0x69, 0x6F)
PINK = RGBColor(0xF9, 0xD5, 0xD7)
BLUSH = RGBColor(0xFC, 0xEE, 0xEF)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x56, 0x5B, 0x57)
MUTED = RGBColor(0x6E, 0x74, 0x7A)
LIGHT = RGBColor(0xF4, 0xF4, 0xF3)
LINE = RGBColor(0xD9, 0xD9, 0xD6)
GREEN = RGBColor(0x1B, 0xAF, 0x7A)
STONE = RGBColor(0x8A, 0x89, 0x84)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
MONO = "Courier New"

prs = Presentation(SRC)
S = list(prs.slides)


# ---------------------------------------------------------------- helpers
def runs_of(text):
    """Split '**bold**' markup into (text, bold) segments."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


def fill_tf(tf, paras, size=14, color=DARK, font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            margins=(0.05, 0.03, 0.05, 0.03), spacing=4):
    """paras: list of str or dict(text, size, color, bold, bullet, align, font, sb)."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(m) for m in margins)
    for i, spec in enumerate(paras):
        spec = {"text": spec} if isinstance(spec, str) else spec
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = spec.get("align", align)
        para.space_before = Pt(spec.get("sb", spacing if i else 0))
        if spec.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.18)))
            for tag, attrs in (("a:buClr", None), ("a:buFont", {"typeface": FONT}), ("a:buChar", {"char": "•"})):
                el = pPr.makeelement(qn(tag), attrs or {})
                if tag == "a:buClr":
                    clr = el.makeelement(qn("a:srgbClr"), {"val": str(spec.get("bucolor", RED))})
                    el.append(clr)
                pPr.append(el)
        for text, bold in runs_of(spec["text"]):
            r = para.add_run()
            r.text = text
            f = r.font
            f.size = Pt(spec.get("size", size))
            f.bold = spec.get("bold", False) or bold
            f.italic = spec.get("italic", False)
            f.name = spec.get("font", font)
            f.color.rgb = spec.get("color", color)
    return tf


def text(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, paras, **kw)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.0):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.text_frame.text = ""
    return sh


def label_box(slide, x, y, w, h, paras, fill, line=None, radius=0.06, **kw):
    sh = box(slide, x, y, w, h, fill, line, radius)
    fill_tf(sh.text_frame, paras, **kw)
    return sh


def card(slide, x, y, w, h, title, body, head=RED, head_h=0.5, title_size=16, body_size=13):
    """Card with a solid header strip and a bordered body."""
    box(slide, x, y, w, h, WHITE, head, radius=0.03, lw=1.25)
    hd = box(slide, x, y, w, head_h, head, None, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    fill_tf(hd.text_frame, [{"text": title, "size": title_size, "bold": True, "color": WHITE}],
            anchor=MSO_ANCHOR.MIDDLE, margins=(0.15, 0.02, 0.1, 0.02))
    text(slide, x + 0.12, y + head_h + 0.08, w - 0.24, h - head_h - 0.14, body, size=body_size)


def picture(slide, name, x, y, w=None, h=None):
    path = f"{IMG}/{name}"
    iw, ih = Image.open(path).size
    if w is not None and h is None:
        h = w * ih / iw
    elif h is not None and w is None:
        w = h * iw / ih
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def line(slide, x1, y1, x2, y2, color=STONE, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    return c


def arrow(slide, x, y, w=0.32, h=0.34, color=ROSE):
    a = box(slide, x, y, w, h, color, None, shape=MSO_SHAPE.RIGHT_ARROW)
    return a


def set_text(shape, value):
    """Replace a text frame's content while keeping the first run's formatting."""
    tf = shape.text_frame
    paras = tf.paragraphs
    first_run = next((r for p in paras for r in p.runs), None)
    rpr = copy.deepcopy(first_run._r.find(qn("a:rPr"))) if first_run is not None else None
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, line_text in enumerate(value.split("\n")):
        para = p0 if i == 0 else tf.add_paragraph()
        if i:
            para._p.insert(0, copy.deepcopy(p0._p.find(qn("a:pPr")))) if p0._p.find(qn("a:pPr")) is not None else None
        r = para.add_run(); r.text = line_text
        if rpr is not None:
            old = r._r.find(qn("a:rPr"))
            if old is not None:
                r._r.remove(old)
            r._r.insert(0, copy.deepcopy(rpr))
    return shape


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in sh.shapes:
                if sub.name == name:
                    return sub
    raise KeyError(name)


def chrome_only(slide, keep_names=()):
    """Keep header group, rule line, title, top-right logo and bottom-right banner; drop the rest."""
    keep = {sh.shape_id for sh in slide.shapes if sh.name in keep_names}
    title = None
    for sh in slide.shapes:
        top, left = sh.top / 914400, sh.left / 914400
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP and top < 0.3 and left < 1:
            keep.add(sh.shape_id)
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE and top < 1:
            keep.add(sh.shape_id)
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE and (left > 12 and top < 0.5 or left > 9.9 and top > 5.9):
            keep.add(sh.shape_id)
        elif sh.has_text_frame and top < 0.3 and 1.4 < left < 1.8 and title is None:
            keep.add(sh.shape_id); title = sh
    for sh in list(slide.shapes):
        if sh.shape_id not in keep:
            sh._element.getparent().remove(sh._element)
    return title


def section(slide, num):
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP and sh.top / 914400 < 0.3:
            for sub in sh.shapes:
                if sub.has_text_frame and re.fullmatch(r"\s*0\d\s*", sub.text_frame.text):
                    set_text(sub, num)


def content_slide(slide, num, title, keep_names=()):
    t = chrome_only(slide, keep_names)
    set_text(t, title)
    t.width = Inches(10.4)
    section(slide, num)


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# ---------------------------------------------------------------- slide 1: cover
set_text(find(S[0], "文本框 6"), "LLM-Driven Hardware Design Automation\nfor Communication Algorithms")
for para_ in find(S[0], "文本框 6").text_frame.paragraphs:
    for r in para_.runs:
        r.font.size = Pt(30)
set_text(find(S[0], "文本框 9"), "Master's Midterm Report · Class of 2024")
notes(S[0], "Midterm report. Topic: using LLMs to design hardware for communication algorithms, "
            "from an ideal formula to verified, synthesizable RTL under system-level quality constraints.")

# ---------------------------------------------------------------- slide 2: contents
for name, value in (("文本框 11", "Research Background and Motivation"),
                    ("文本框 17", "Research Problems and Framework"),
                    ("文本框 20", "Methods and Preliminary Results"),
                    ("文本框 23", "Open Issues and Schedule")):
    set_text(find(S[1], name), value)

# ---------------------------------------------------------------- slide 3: background
s = S[2]
content_slide(s, "01", "Research Background", keep_names=("图片 33",))
flow = find(s, "图片 33")
flow.height = Inches(5.3); flow.width = Inches(5.3 * 4.65 / 5.43)
flow.left = Inches(8.55); flow.top = Inches(1.0)
text(s, 0.5, 1.0, 7.9, 0.5, [{"text": "Communication baseband: from formula to fixed-point RTL", "size": 18,
                               "bold": True, "color": RED}])
facts = [("Where the cost is", "Digital front end, FFT, equalization and demodulation dominate baseband area "
                                "and power."),
         ("What is fixed, what is free", "The ideal formula is fixed; structure, fixed-point format and "
                                         "micro-architecture are free."),
         ("How designs are accepted", "By system metrics (SNR, EVM, BER) within a loss budget, not by bit-exactness "
                                      "to the formula, so approximation is legal."),
         ("How it is done today", "Every arrow on the right is manual iteration, repeated whenever the spec changes.")]
y = 1.55
for head, body in facts:
    box(s, 0.5, y + 0.06, 0.07, 0.68, RED, None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.7, y, 7.7, 0.85, [{"text": head, "size": 15, "bold": True, "color": RED},
                                {"text": body, "size": 14, "sb": 1}])
    y += 0.86
text(s, 0.5, 5.0, 7.9, 0.4, [{"text": "Existing automation cannot use this legal approximation:", "size": 15,
                               "bold": True}])
for i, (head, body) in enumerate((("Rule-based", "exact transformations only"),
                                  ("LLM-based", "pass / fail checking only"))):
    label_box(s, 0.5 + i * 4.0, 5.5, 3.85, 0.7, [{"text": f"**{head}:** {body}", "size": 14, "color": RED}],
              BLUSH, anchor=MSO_ANCHOR.MIDDLE, margins=(0.15, 0.03, 0.1, 0.03))
notes(s, "Our object is the numeric datapath of the communication baseband: front end, FFT, equalization and "
         "demodulation, which dominate area and power. The ideal formula is fixed, but the structure, fixed-point "
         "format and micro-architecture are free. Designs are accepted by system metrics within a loss budget, "
         "not by bit-exactness, so approximation is legal. Today each step from formula to RTL is manual "
         "iteration. Two families of automation exist, and neither can use this legal approximation: rule-based "
         "methods only apply exact transformations, and LLM-based methods only check whether a design passes.")

# ---------------------------------------------------------------- slide 4: rule-based
s = S[3]
content_slide(s, "01", "Existing Work: Rule-Based", keep_names=("图片 1",))
fig = find(s, "图片 1")
fig.left = Inches(0.4); fig.top = Inches(1.35); fig.width = Inches(6.2); fig.height = Inches(6.2 * 2.25 / 6.41)
text(s, 0.4, 3.62, 6.2, 0.4, [{"text": "SPIRAL: search over expert rewrite rules for linear transforms.",
                               "size": 11, "color": MUTED}])
rules = [("Exact by construction", "Rewrite rules, HLS and e-graph optimization preserve the semantics; "
                                   "FloPoCo bounds error per operator. No error is allowed at the system level."),
         ("Structure is given", "Word-length optimization assigns bits on a fixed dataflow graph; the designer "
                                "still chooses the structure."),
         ("Bounded by the rule library", "Only transformations encoded by experts are reachable; cross-operator "
                                         "restructuring is hard to express.")]
y = 1.15
for head, body in rules:
    box(s, 6.95, y + 0.06, 0.07, 1.05, RED, None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 7.15, y, 5.65, 1.25, [{"text": head, "size": 16, "bold": True, "color": RED},
                                  {"text": body, "size": 14, "sb": 2}])
    y += 1.4
label_box(s, 0.4, 5.3, 12.4, 0.62,
          [{"text": "Missed: legal approximations with the largest savings — CIC for an ideal low-pass, truncated NCO "
                    "phase, fewer filter coefficient bits.", "size": 14, "bold": True, "color": RED}],
          BLUSH, anchor=MSO_ANCHOR.MIDDLE, margins=(0.2, 0.03, 0.15, 0.03))
text(s, 0.4, 6.05, 9.5, 0.4, [{"text": "SPIRAL (Püschel 2005); FloPoCo (de Dinechin 2011); word-length optimization "
                                        "(Constantinides 2003); ROVER (Coward 2024); SkyEgg (Zou 2025); Vitis HLS.",
                                "size": 9.5, "color": MUTED}])
notes(s, "Rule-based methods include DSP generators such as SPIRAL, HLS, e-graph rewriting and word-length "
         "optimization. Their common property is that transformations are exact: rewrite rules and HLS preserve "
         "the program semantics, and FloPoCo bounds error only per operator. Word-length optimization tunes bits but "
         "keeps the dataflow structure given by the designer, and every reachable transformation must be encoded in "
         "a rule library. As a result, the approximations that save the most area, such as replacing an ideal "
         "low-pass with a CIC, are outside their search space.")

# ---------------------------------------------------------------- slide 5: LLM-based
s = S[4]
content_slide(s, "01", "Existing Work: LLM-Based", keep_names=("图片 2",))
fig = find(s, "图片 2")
fig.left = Inches(0.4); fig.top = Inches(1.2); fig.width = Inches(5.2); fig.height = Inches(5.2 * 2.33 / 6.41)
text(s, 0.4, 3.12, 5.2, 0.4, [{"text": "Typical loop: the LLM writes code, tools test it, failures trigger a retry.",
                               "size": 11, "color": MUTED}])
llm = [("Only pass / fail", "Functional tests or equivalence checks decide acceptance; designs must match the "
                            "reference, so no approximation is allowed."),
       ("Coarse feedback", "A binary verdict or a single score: no gap per constraint, no location of the error."),
       ("Coarse edits", "Whole Verilog or HLS C files are rewritten; structure, bit-width and rounding cannot be "
                        "changed one at a time.")]
y = 1.1
for head, body in llm:
    box(s, 5.95, y + 0.06, 0.07, 0.62, RED, None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 6.15, y, 6.65, 0.8, [{"text": head, "size": 15, "bold": True, "color": RED},
                                 {"text": body, "size": 13, "sb": 1}])
    y += 0.8
rows = [("", "Allowed error", "Acceptance signal", "Edit granularity"),
        ("Rule-based", "None (exact, or per-operator bound)", "Rule applicability", "Expert rewrite rules"),
        ("LLM-based", "None (functional equivalence)", "Pass / fail or one score", "Whole file"),
        ("This work", "System tolerance (SNR / EVM / BER)", "Per-constraint gap + real PPA",
         "Typed primitive edits (REWRITE / TUNE)")]
tbl = s.shapes.add_table(4, 4, Inches(0.4), Inches(3.65), Inches(12.4), Inches(0.46 * 4)).table
for j, wd in enumerate((1.7, 3.8, 3.2, 3.7)):
    tbl.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED if i == 0 else (BLUSH if i == 3 else (WHITE if i % 2 else LIGHT))
        fill_tf(cell.text_frame, [{"text": val, "size": 13, "bold": i in (0, 3) or j == 0,
                                   "color": WHITE if i == 0 else (RED if i == 3 else DARK)}],
                anchor=MSO_ANCHOR.MIDDLE, margins=(0.1, 0.02, 0.08, 0.02))
text(s, 0.4, 5.65, 9.5, 0.5, [{"text": "RTLCoder, AutoChip, VerilogCoder; EvolVE, REvolution, COEVO; CktFormalizer; "
                                        "C-to-HLS refactoring (arXiv:2412.00214).", "size": 9.5, "color": MUTED}])
notes(s, "LLM-based work covers fine-tuned RTL models, agent loops that repair Verilog from compiler and simulation "
         "feedback, and evolutionary search that optimizes PPA. Three limits are shared. Acceptance is pass or fail "
         "against a reference, so approximation is not allowed. Feedback is a binary verdict or a single score, so "
         "the model does not know how far each constraint is from passing or where the error comes from. And the "
         "edit unit is a whole Verilog or HLS file, so structure, bit-width and rounding cannot be changed one at "
         "a time. The table summarizes the difference: this work allows error within the system tolerance, returns "
         "the gap of every constraint plus real PPA, and edits typed primitives with REWRITE and TUNE.")

# ---------------------------------------------------------------- slide 6: motivation
s = S[5]
content_slide(s, "01", "Motivation: Structure Dominates Cost")
picture(s, "motivation_en.png", 0.5, 1.1, w=6.3)
text(s, 0.5, 5.05, 6.3, 0.35, [{"text": "Same DDC formula and system metric; real synthesis area on Nangate45.",
                                 "size": 11, "color": MUTED}])
text(s, 7.2, 1.1, 5.6, 0.5, [{"text": "Three properties of communication datapaths", "size": 18, "bold": True,
                               "color": RED}])
items = [("Fixed math, free implementation",
          "Structure, fixed-point format and micro-architecture can all change; area differs by more than 10×."),
         ("Approximation is legal",
          "Acceptance follows the system loss budget (SNR, EVM, ACLR), not bit-exactness to the ideal algorithm."),
         ("Operators are coupled",
          "Error injected at one stage is filtered or amplified downstream; the budget must be allocated over the chain.")]
for i, (head, body) in enumerate(items):
    y = 1.75 + i * 1.2
    label_box(s, 7.2, y, 0.5, 0.5, [{"text": str(i + 1), "size": 18, "bold": True, "color": WHITE,
                                      "align": PP_ALIGN.CENTER}], RED, radius=0.5, anchor=MSO_ANCHOR.MIDDLE,
              margins=(0, 0, 0, 0))
    text(s, 7.85, y - 0.05, 4.95, 1.1, [{"text": head, "size": 15, "bold": True},
                                        {"text": body, "size": 13, "color": GRAY, "sb": 2}])
label_box(s, 0.5, 5.35, 12.3, 0.6,
          [{"text": "The 8–46× gap comes from structural transformation, not bit-width tuning — "
                    "and structure choice is the least automated step today.", "size": 14, "bold": True,
            "color": RED}], BLUSH, anchor=MSO_ANCHOR.MIDDLE, margins=(0.2, 0.03, 0.15, 0.03))
notes(s, "Take the digital down-converter, which almost every receiver contains. Mapping the formula term by term "
         "gives a full-period lookup table and one long FIR. The textbook multirate structure uses NCO, CIC, "
         "half-band and compensation filters. Under the same system metric the direct design is 8 to 46 times "
         "larger. The gap is structural. Three properties make this possible: the math is fixed but the "
         "implementation is free, approximation is legal under the system loss budget, and operators are "
         "coupled, so the error budget must be allocated over the whole chain.")

# ---------------------------------------------------------------- slide 7: research problems
s = S[6]
content_slide(s, "02", "Research Problems")
text(s, 0.5, 1.0, 12.3, 0.5, [{"text": "**Route:** the LLM proposes design edits; verification and synthesis tools "
                                        "decide what is accepted.", "size": 16}])
probs = [("Q1  Representation", "Which representation should the LLM edit?",
          ["Verilog hides structure in enables and shift registers; ~6 min per evaluation",
           "Formula edits cannot be mapped to hardware automatically",
           "Need: typed, executable, bit-accurate, and able to constrain the RTL"]),
         ("Q2  Judgement", "When is an edited design acceptable?",
          ["Candidate vs. ideal: approximate, judged end-to-end by system metrics with full interference coverage",
           "RTL vs. candidate: must be bit- and transaction-exact",
           "Two relations, two different criteria"]),
         ("Q3  Search", "How should many edits be organized?",
          ["One-shot generation: only 9 of 80 full DDC designs met the spec",
           "Keep near-feasible designs as stepping stones",
           "Never overwrite feasible designs; trade off area, timing and power"])]
for i, (head, q, bullets) in enumerate(probs):
    x = 0.5 + i * 4.25
    card(s, x, 1.65, 3.85, 3.6, head,
         [{"text": q, "size": 15, "bold": True, "italic": True, "color": RED}] +
         [{"text": b, "size": 15, "bullet": True, "sb": 8} for b in bullets])
    if i < 2:
        arrow(s, x + 3.92, 3.45)
text(s, 0.5, 5.6, 9.3, 0.45, [{"text": "Represent → Judge → Search: each question presupposes the previous one.",
                                 "size": 15, "bold": True, "color": RED}])
notes(s, "Our route: the LLM only proposes edits; tools decide acceptance. Three questions follow, in order. "
         "First, the representation: Verilog is too low level and slow to evaluate, and a formula cannot be mapped "
         "to hardware automatically, so we need a typed, executable middle layer. Second, judgement: there are two "
         "different relations. The candidate only needs to approximate the ideal function within the system "
         "tolerance, while the RTL must be exactly faithful to the candidate. Third, search: one-shot generation "
         "rarely works, so we need to keep near-feasible designs and never overwrite feasible ones. Each question "
         "builds on the previous one.")

# ---------------------------------------------------------------- slide 8: framework
s = S[7]
content_slide(s, "02", "Framework Overview")
picture(s, "fig_framework_v5.png", 0.5, 0.9, w=12.3)
text(s, 0.5, 6.6, 12.3, 0.45, [{"text": "Candidate ≈ ideal function within the deployment contract;  "
                                         "RTL ≡ the frozen candidate.", "size": 15, "bold": True, "color": RED}])
notes(s, "The framework has five stages in execution order. Stage 0 turns the task contract into an initial "
         "primitive program, the root of the search tree. Stage 1: the design agent selects a parent from the "
         "AND/OR tree and applies REWRITE or TUNE to produce a complete child program. Stage 2 checks whether the "
         "candidate is deployable: type and legality checks, verification of the declared relation, and the "
         "system-quality check. Bit-exact edits inherit the parent's quality verdict. Failed candidates return as "
         "BROKEN or OPEN with their quality gap. Stage 3 generates RTL inside the frozen contract and checks "
         "fidelity, with bounded local repair. Stage 4 measures real PPA, feeds it back to the tree, and runs an "
         "independent final check on the Pareto front. The final check never feeds back into search.")

# ---------------------------------------------------------------- slide 9: method 1
s = S[8]
content_slide(s, "03", "Method 1: Typed Primitive Representation")
code = ["# primitive program (link A, I path, simplified)",
        "x  = g.input()             # (12,11,s,d0,0)",
        "ph = g.counter(fcw, 32)    # phase accumulator",
        "a  = g.bits(ph, 31, 24)    # 8-bit address",
        "c  = g.lut(a, COS256, 12, 11)",
        "m  = g.mul(x, c)           # (24,22,s,d0,0)",
        "i  = g.q(m, 16, 15, 'floor', 'wrap')",
        "i  = g.integrate(i, 24)    # accumulator",
        "i  = g.decimate(i, 32)     # -> rate domain d1",
        "g.output(g.q(i, 16, 15), ...)"]
label_box(s, 0.5, 1.1, 5.7, 2.7,
          [{"text": c, "size": 13, "font": MONO, "color": GRAY if c.startswith("#") else DARK, "sb": 1}
           for c in code], LIGHT, LINE, radius=0.03, margins=(0.18, 0.15, 0.1, 0.1))
text(s, 0.5, 3.9, 5.7, 0.4, [{"text": "Type per signal τ = (w, f, σ, d, ℓ): width, fraction bits, sign, rate "
                                       "domain, latency.", "size": 12, "color": MUTED}])
for i, (big, small) in enumerate((("1–3 s", "per chain evaluation on the integer model (~6 min for RTL sim + synthesis)"),
                                  ("0", "RTL / model mismatches on all reported designs"))):
    x = 0.5 + i * 2.95
    label_box(s, x, 4.55, 2.75, 1.15,
              [{"text": big, "size": 24, "bold": True, "color": RED},
               {"text": small, "size": 10.5, "color": GRAY, "sb": 0}], BLUSH, radius=0.05,
              anchor=MSO_ANCHOR.TOP, margins=(0.15, 0.1, 0.1, 0.05))
blocks = [("Edit target", "A primitive program — neither the formula nor Verilog. 16 operator-agnostic primitives; "
                          "no FIR / CIC / FFT blocks, which would freeze the structure."),
          ("Checks by construction", "Latency alignment, rate-domain checks and valid-signal derivation; violations "
                                     "return classified construction diagnostics."),
          ("Edit = action × relation", "REWRITE (decomposition, topology, state, rate) or TUNE (width, rounding, "
                                       "pipelining, sharing), each declared bit-exact, math-equivalent or approximate."),
          ("Frozen implementation contract", "After acceptance, interface, quantization, state, latency / II and "
                                             "resources are fixed for RTL generation. The RTL generator itself is not "
                                             "claimed; all reported runs use a deterministic compiler.")]
y = 1.1
for head, body in blocks:
    box(s, 6.6, y + 0.06, 0.07, 1.0, RED, None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 6.8, y, 6.0, 1.15, [{"text": head, "size": 15, "bold": True, "color": RED},
                                {"text": body, "size": 13, "color": DARK, "sb": 2}])
    y += 1.22
notes(s, "Method 1 is the representation. The LLM edits a primitive program, not the formula and not Verilog. "
         "There are 16 operator-agnostic primitives such as add, mul, q for quantization, lut, delay, integrate, "
         "decimate and mac. There are deliberately no FIR or CIC blocks, because high-level blocks would freeze the "
         "structure. Every signal carries a type: width, fraction bits, sign, rate domain and latency, so latency "
         "alignment, rate checks and valid signals are derived automatically. Each edit is described by two "
         "orthogonal attributes: the action, REWRITE or TUNE, and the declared semantic relation to the parent. "
         "Once a candidate is accepted, its interface, quantization, state, latency and resources become a frozen "
         "implementation contract for RTL generation. The RTL generator is not a contribution; all reported runs "
         "use our deterministic compiler.")

# ---------------------------------------------------------------- slide 10: method 2
s = S[9]
content_slide(s, "03", "Method 2: Two Trust Relations")
card(s, 0.5, 1.1, 6.0, 2.9, "2a   Candidate ≈ Ideal   (system tolerance)",
     [{"text": "∀ s ∈ S :   M( ⟦p⟧(xₛ), F(x̃ₛ) ) ∈ T", "size": 16, "font": "Cambria Math", "color": RED,
       "align": PP_ALIGN.CENTER},
      {"text": "Judged once at the chain output; no per-operator specs", "bullet": True, "sb": 10},
      {"text": "Accept if implementation error power ≤ p × allowed error power", "bullet": True, "sb": 6},
      {"text": "Linear parts: full-band sweep; others: search scenarios, then held-out and statistical "
               "re-check with a failure-probability bound", "bullet": True, "sb": 6}], body_size=14)
card(s, 6.8, 1.1, 6.0, 2.9, "2b   RTL ≡ Candidate   (implementation fidelity)",
     [{"text": "∀ x, t :   RTLₚ(x, t) = ⟦p⟧(x, t)", "size": 16, "font": "Cambria Math", "color": GRAY,
       "align": PP_ALIGN.CENTER},
      {"text": "Same output every cycle, plus transactions, reset, latency and II", "bullet": True, "sb": 10,
       "bucolor": GRAY},
      {"text": "Ensures the judged object is the hardware that ships", "bullet": True, "sb": 6, "bucolor": GRAY},
      {"text": "Currently RTL co-simulation; formal equivalence checking planned", "bullet": True, "sb": 6,
       "bucolor": GRAY}], head=GRAY, body_size=14)
label_box(s, 0.5, 4.25, 12.3, 0.8,
          [{"text": "**Claim check and inheritance.**  Bit-exact claims are compared sample by sample with the parent "
                    "and inherit its quality verdict; math-equivalent claims are compared with widened quantizers; "
                    "failed claims are judged as approximate.", "size": 14}],
          BLUSH, anchor=MSO_ANCHOR.MIDDLE, margins=(0.2, 0.03, 0.15, 0.03))
text(s, 0.5, 5.25, 9.4, 0.7,
     [{"text": "**Coverage lesson.**  With 3 blocker tones per scenario, only 11 of 77 best designs passed held-out "
               "scenarios (worst −34 dB). Interference must cover all aliasing bands.", "size": 14, "color": GRAY}])
notes(s, "Method 2 separates two trust relations. 2a: the candidate only needs to approximate the ideal formula "
         "within the deployment tolerance. It is judged once at the chain output with system metrics; the "
         "acceptance rule is that implementation error power must not exceed p times the allowed error power. "
         "Linear parts are checked with a full-band sweep, the rest with search scenarios, then held-out and "
         "statistical re-checks. 2b: the RTL must be exactly faithful to the candidate, cycle by cycle, including "
         "transactions, reset, latency and II. These two relations cannot be merged into one equivalence claim. "
         "Before judging quality we verify the declared relation; verified bit-exact edits inherit the parent's "
         "verdict. One lesson from development: when scenarios covered only a few blocker tones, search simply "
         "built notches at those frequencies, and most designs failed held-out scenarios.")

# ---------------------------------------------------------------- slide 11: method 3
s = S[10]
content_slide(s, "03", "Method 3: AND/OR Tree Search")
label_box(s, 1.9, 1.1, 3.0, 0.55, [{"text": "Root: initial primitive program", "size": 12, "bold": True,
                                     "align": PP_ALIGN.CENTER}], WHITE, STONE, anchor=MSO_ANCHOR.MIDDLE)
label_box(s, 0.5, 2.05, 2.75, 0.62, [{"text": "OR node", "size": 13, "bold": True, "color": RED,
                                      "align": PP_ALIGN.CENTER},
                                     {"text": "region + boundary contract", "size": 11, "align": PP_ALIGN.CENTER,
                                      "sb": 0}], BLUSH, RED, anchor=MSO_ANCHOR.MIDDLE)
label_box(s, 3.55, 2.05, 2.75, 0.62, [{"text": "AND node", "size": 13, "bold": True, "color": GRAY,
                                       "align": PP_ALIGN.CENTER},
                                      {"text": "one decomposition", "size": 11, "align": PP_ALIGN.CENTER,
                                       "sb": 0}], LIGHT, GRAY, anchor=MSO_ANCHOR.MIDDLE)
line(s, 3.4, 1.65, 1.875, 2.05); line(s, 3.4, 1.65, 4.925, 2.05)
for i, lab in enumerate(("REWRITE", "TUNE")):
    label_box(s, 0.5 + i * 1.45, 3.0, 1.3, 0.42, [{"text": lab, "size": 11, "bold": True, "color": RED,
                                                   "align": PP_ALIGN.CENTER}], WHITE, RED, anchor=MSO_ANCHOR.MIDDLE)
    line(s, 1.875, 2.67, 1.15 + i * 1.45, 3.0)
for i, lab in enumerate(("sub-region 1", "sub-region 2")):
    label_box(s, 3.55 + i * 1.45, 3.0, 1.3, 0.42, [{"text": lab, "size": 11, "bold": True, "color": GRAY,
                                                    "align": PP_ALIGN.CENTER}], WHITE, GRAY, anchor=MSO_ANCHOR.MIDDLE)
    line(s, 4.925, 2.67, 4.2 + i * 1.45, 3.0)
text(s, 0.5, 3.45, 2.75, 0.35, [{"text": "alternatives", "size": 11, "italic": True, "color": MUTED,
                                 "align": PP_ALIGN.CENTER}])
text(s, 3.55, 3.45, 2.75, 0.35, [{"text": "implemented jointly", "size": 11, "italic": True, "color": MUTED,
                                  "align": PP_ALIGN.CENTER}])
for i, (lab, col, desc) in enumerate((("BROKEN", STONE, "construction failed"), ("OPEN", ROSE, "legal; gap v > 0"),
                                     ("FEASIBLE", GREEN, "constraints met with margin"),
                                     ("VERIFIED", RED, "also passed the final check"))):
    text(s, 0.45 + i * 1.47, 4.4, 1.45, 0.6, [{"text": desc, "size": 10.5, "color": GRAY, "align": PP_ALIGN.CENTER}])
    label_box(s, 0.5 + i * 1.47, 3.95, 1.35, 0.4, [{"text": lab, "size": 11, "bold": True, "color": WHITE,
                                                    "align": PP_ALIGN.CENTER}], col, anchor=MSO_ANCHOR.MIDDLE,
              radius=0.3, margins=(0, 0, 0, 0))
steps = [("Select", "Small-gap OPEN node → TUNE; stalled TUNE or a structural bottleneck → REWRITE"),
         ("Expand", "The LLM writes a complete child program; parents are never overwritten"),
         ("Judge", "Method 2 at the chain root; local metrics are never summed"),
         ("Back up", "Quality gap and measured PPA are written back to the node"),
         ("Output", "Pareto front of FEASIBLE nodes → independent final check (no feedback) → VERIFIED")]
for i, (head, body) in enumerate(steps):
    y = 1.1 + i * 0.83
    label_box(s, 6.8, y + 0.05, 0.45, 0.45, [{"text": str(i + 1), "size": 14, "bold": True, "color": WHITE,
                                              "align": PP_ALIGN.CENTER}], RED, radius=0.5,
              anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    text(s, 7.4, y - 0.02, 5.4, 0.82, [{"text": head, "size": 14, "bold": True, "color": RED},
                                       {"text": body, "size": 12, "sb": 0}])
label_box(s, 0.5, 5.25, 12.3, 0.72,
          [{"text": "**Status.**  OR branches, node states, parent retention, root judgement and back-up are "
                    "implemented; AND sub-frontier composition is pending (protocol prototype). Reported runs use "
                    "degenerate forms: a serial agent (greedy OR chain) and an OpenEvolve population (flat).",
            "size": 12}], LIGHT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.2, 0.03, 0.15, 0.03))
notes(s, "Method 3 organizes the search as an AND/OR tree whose root is the initial primitive program. An OR node "
         "is a region with a boundary contract; different REWRITE or TUNE children are alternatives. An AND node is "
         "a decomposition whose sub-regions must be implemented together. Every combination is materialized as a "
         "complete program and judged at the chain root, so local metrics are never simply added up. Nodes are "
         "append-only, with four states: BROKEN, OPEN, FEASIBLE and VERIFIED. Each round selects a node, expands it, "
         "judges the child and writes back the quality gap and measured PPA. The output is the Pareto front, "
         "checked independently. To be precise about status: OR branches and states are implemented, but AND "
         "sub-frontier composition is still pending, and our reported runs use degenerate forms of this tree.")

# ---------------------------------------------------------------- slide 12: results
s = S[11]
content_slide(s, "03", "Results: Link A, R32 Narrowband")
rows = [("Design", "Area (µm²)", "Fmax (MHz)", "Power (mW)", "Held-out"),
        ("Direct mapping", "715,916", "304", "182.0", "Pass"),
        ("Evolutionary search", "93,689", "104", "32.0", "Pass"),
        ("Template param. search", "27,345", "250", "9.66", "Pass"),
        ("DeepSeek agent", "110,467", "244", "30.7", "Pending"),
        ("Our agent", "14,639", "423", "4.78", "Pass")]
tbl = s.shapes.add_table(len(rows), 5, Inches(0.5), Inches(1.15), Inches(6.2), Inches(0.42 * len(rows))).table
for j, wd in enumerate((2.1, 1.1, 1.0, 1.0, 1.0)):
    tbl.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED if i == 0 else (BLUSH if i == len(rows) - 1 else (WHITE if i % 2 else LIGHT))
        fill_tf(cell.text_frame, [{"text": val, "size": 12, "bold": i in (0, len(rows) - 1),
                                   "color": WHITE if i == 0 else (RED if i == len(rows) - 1 else DARK),
                                   "align": PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT}],
                anchor=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.02, 0.08, 0.02))
for i, (big, small) in enumerate((("2.0%", "of direct-mapping area"), ("−46%", "area vs. template search"),
                                  ("+69% / −51%", "Fmax / power vs. template search"))):
    x = 0.5 + i * 2.1
    label_box(s, x, 3.85, 1.95, 1.15, [{"text": big, "size": 22 if i < 2 else 17, "bold": True, "color": RED},
                                       {"text": small, "size": 11, "color": GRAY, "sb": 2}], BLUSH,
              anchor=MSO_ANCHOR.TOP, margins=(0.12, 0.15, 0.08, 0.03))
text(s, 0.5, 5.15, 6.2, 0.75, [{"text": "45 dB blockers, SNR loss ≤ 1.5 dB. Yosys + ABC on Nangate45; OpenSTA at "
                                          "2 ns, toggle rate 0.1. Our agent used 41 tool evaluations.",
                                 "size": 10.5, "color": MUTED}])
picture(s, "agent_trace_en.png", 6.95, 1.15, w=5.85)
text(s, 6.95, 4.45, 5.85, 0.9, [{"text": "Agent trace: 30 candidates. Infeasible intermediate designs are kept "
                                          "and repaired; final search margin 0.38 dB, held-out margin 0.55 dB. "
                                          "Points use the proxy area model; dashed lines are real synthesis areas "
                                          "for reference only.",
                                  "size": 12, "color": GRAY}])
notes(s, "Main result on link A, R32 narrowband, with 45 dB blockers. All numbers are real synthesis results under "
         "the same judge and the same flow. Our agent reaches 14.6 thousand square microns, 2 percent of the direct "
         "mapping, and compared with a template parameter search with 4000 evaluations it saves 46 percent area, "
         "raises Fmax by 69 percent and cuts power by 51 percent, using only 41 tool evaluations. The trace on the "
         "right shows that the first feasible design came after several infeasible attempts, which were kept and "
         "repaired with diagnostics. The remaining specs are being measured.")

# ---------------------------------------------------------------- slide 13: ablations
s = S[12]
content_slide(s, "03", "Ablations: One Component at a Time")
text(s, 0.5, 1.0, 12.3, 0.45, [{"text": "Each ablation replaces one component while keeping the model, budget and "
                                         "judge fixed.", "size": 15}])
abl = [("Method 1", "1.2–5.1×", "LLM edits Verilog directly",
        "Area vs. primitive representation on valid specs; on wideband specs the start point alone exceeds 900 s of "
        "synthesis and no feasible design is found."),
       ("Method 2", "36–42%", "Stage-wise acceptance",
        "Area saved by chain-level over stage-wise acceptance on narrowband specs; on wideband specs stage-wise "
        "design has no feasible solution."),
       ("Method 3", "≈ 7×", "Random search instead of the LLM",
        "Geometric-mean area ratio at 120 evaluations: 0.871 / 0.884 for random edits vs. 0.131 for the LLM."),
       ("Method 3", "−40%", "Metric feedback vs. a single score",
        "11 of 12 pairs better (Wilcoxon p = 0.001); adding diagnosis gives no further gain (p = 0.52).")]
for k, (tag, big, head, body) in enumerate(abl):
    x = 0.5 + (k % 2) * 6.25
    y = 1.6 + (k // 2) * 2.05
    box(s, x, y, 6.05, 1.85, WHITE, LINE, radius=0.04)
    text(s, x + 0.15, y + 0.3, 2.1, 1.0, [{"text": big, "size": 30, "bold": True, "color": RED}],
         anchor=MSO_ANCHOR.MIDDLE)
    label_box(s, x + 0.2, y + 1.3, 1.2, 0.34, [{"text": tag, "size": 10.5, "bold": True, "color": WHITE,
                                                "align": PP_ALIGN.CENTER}], GRAY, radius=0.3,
              anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    text(s, x + 2.35, y + 0.12, 3.55, 1.65, [{"text": head, "size": 14, "bold": True},
                                             {"text": body, "size": 12, "color": GRAY, "sb": 3}])
text(s, 0.5, 5.75, 9.4, 0.6, [{"text": "Evolutionary search on Link A: −71.5% to −89.0% area on 6 development specs, "
                                        "−40.5% to −92.5% on 6 held-out specs, 0 RTL mismatches.", "size": 12.5,
                                "color": DARK}])
notes(s, "Each ablation replaces one component with everything else fixed. Method 1: if the LLM edits Verilog "
         "directly, areas are 1.2 to 5.1 times larger and wideband specs cannot even finish. Method 2: replacing "
         "chain-level acceptance with stage-wise budgets costs 36 to 42 percent area on narrowband specs and has no "
         "feasible solution on wideband specs. Method 3: random search instead of the LLM ends about 7 times larger, "
         "and returning only a single score instead of all metric values costs about 40 percent area.")

# ---------------------------------------------------------------- slide 14: benchmark and link B
s = S[13]
content_slide(s, "03", "Benchmark and Second Chain")
chains = [("ID", "Chain", "Output metric", "Status"),
          ("A", "Receive DDC", "In-band SNR, ripple", "6 dev + 6 held-out"),
          ("B", "Transmit DUC", "EVM, ACLR", "3 specs"),
          ("C", "OFDM transmit", "EVM, out-of-band leakage", "Planned"),
          ("D", "OFDM receive inner loop", "EVM, BER loss", "Planned"),
          ("E", "Burst sync front end", "Detection prob., CFO error", "Planned"),
          ("F", "Polyphase channelizer", "Sub-band leakage, EVM", "Planned"),
          ("G", "WCDMA Rake branch", "Combined SNR", "Planned"),
          ("H", "Fractional resampler", "EVM, image rejection", "Planned")]
tbl = s.shapes.add_table(len(chains), 4, Inches(0.5), Inches(1.1), Inches(7.4), Inches(0.4 * len(chains))).table
for j, wd in enumerate((0.5, 2.5, 2.6, 1.8)):
    tbl.columns[j].width = Inches(wd)
for i, row in enumerate(chains):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        active = i in (1, 2)
        cell.fill.fore_color.rgb = RED if i == 0 else (BLUSH if active else (WHITE if i % 2 else LIGHT))
        fill_tf(cell.text_frame, [{"text": val, "size": 12, "bold": i == 0 or (active and j < 2),
                                   "color": WHITE if i == 0 else (RED if active else DARK)}],
                anchor=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.02, 0.06, 0.02))
text(s, 0.5, 5.0, 7.4, 0.7, [{"text": "Each chain is accepted only at its output. Operator tasks: 7 families "
                                        "(complex multiply, sin/cos, NCO, FIR, atan2, LLR, matched filter).",
                                "size": 12, "color": GRAY}])
card(s, 8.3, 1.1, 4.5, 2.5, "Link B: transmit DUC",
     [{"text": "−45% to −50% area", "size": 22, "bold": True, "color": RED},
      {"text": "Added with a 20-line spec file; no system changes", "bullet": True, "sb": 6, "size": 13},
      {"text": "Auto-generated start point meets all 3 specs", "bullet": True, "sb": 4, "size": 13},
      {"text": "0 RTL / model mismatches", "bullet": True, "sb": 4, "size": 13}])
card(s, 8.3, 3.8, 4.5, 1.7, "Open-source baseline (USRP)",
     [{"text": "All 12 specs fail under the dense-blocker judge (−8.9 to −70 dB); re-designing the half-band taps "
               "still fails (−8 to −28 dB).", "size": 13}], head=GRAY, head_h=0.45, title_size=14)
notes(s, "The benchmark uses formula, metric and tolerance as the specification, and accepts each chain only at its "
         "output. Links A and B are built; the others are planned and selected by computation pattern. Link B, the "
         "transmit DUC, was added with a 20-line spec file and no system change; evolutionary search reduced area "
         "by 45 to 50 percent with zero RTL mismatches. As a reference, the open-source USRP chain fails all 12 "
         "specs under our dense-blocker judge, which shows that fixed generic structures do not fit these specs.")

# ---------------------------------------------------------------- slide 15: open issues
s = S[14]
content_slide(s, "04", "Open Issues")
text(s, 0.5, 1.0, 12.3, 0.45, [{"text": "Agent results currently cover one spec (R32 narrowband); the next step is "
                                         "all specs with multiple seeds.", "size": 15}])
issues = [("AND/OR composition", "Implement sub-frontier composition and strict AND back-up; compare at equal "
                                 "budget with best-of-N, a flat agent and an OR-only tree."),
          ("Final-check isolation", "In the first tree-search round, held-out slack leaked to the agent. Isolate "
                                    "the final check, freeze the feedback protocol and rerun."),
          ("RTL coverage", "The compiler lacks reset-less data registers and pipelined shared multipliers; a "
                           "lowering agent will be compared on coverage, fidelity pass rate, repair cost and PPA."),
          ("Formal checks", "Claim checks and fidelity are simulation-based; add formal equivalence, SAT or "
                            "exhaustive checks for stateless exact rewrites, and coverage criteria for more metrics.")]
for k, (head, body) in enumerate(issues):
    x = 0.5 + (k % 2) * 6.25
    y = 1.6 + (k // 2) * 2.05
    box(s, x, y, 6.05, 1.85, WHITE, LINE, radius=0.04)
    box(s, x, y, 0.09, 1.85, RED, None, shape=MSO_SHAPE.RECTANGLE)
    label_box(s, x + 0.3, y + 0.25, 0.5, 0.5, [{"text": str(k + 1), "size": 16, "bold": True, "color": WHITE,
                                                "align": PP_ALIGN.CENTER}], RED, radius=0.5,
              anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    text(s, x + 1.0, y + 0.2, 4.9, 1.55, [{"text": head, "size": 16, "bold": True, "color": RED},
                                          {"text": body, "size": 14, "color": DARK, "sb": 4}])
text(s, 0.5, 5.75, 9.4, 0.45, [{"text": "Until issues 1–2 are closed, agent results show feasibility, not "
                                         "search superiority.", "size": 14, "bold": True, "color": RED}])
notes(s, "Four open issues. First, AND/OR composition is not complete, so the search contribution must still be "
         "confirmed against best-of-N, a flat agent and an OR-only tree at equal budget. Second, in the first round "
         "of tree-search experiments the held-out slack leaked to the agent; we will isolate the final check and "
         "rerun. Third, the deterministic compiler does not cover every structure, so we will evaluate a lowering "
         "agent inside the frozen contract. Fourth, claim checks and fidelity are currently simulation-based, and "
         "we plan formal checks.")

# ---------------------------------------------------------------- slide 17: schedule
s = S[16]
slots = [("矩形 54", "文本框 56", "26.10–26.11", "Isolate final check; AND/OR composition; agent on all specs"),
         ("矩形 64", "文本框 66", "26.12–27.01", "OFDM RX and sync chains; lowering agent; formal checks; baselines"),
         ("矩形 67", "文本框 69", "27.02–27.03", "Main experiments; held-out re-check on all chains"),
         ("矩形 79", "文本框 3", "27.03–27.04", "Write thesis; submit paper"),
         ("矩形 76", "文本框 78", "27.05 (est.)", "Thesis defense")]
for head, body, date, desc in slots:
    set_text(find(s, head), date)
    tb = set_text(find(s, body), desc)
    for p_ in tb.text_frame.paragraphs:
        for r in p_.runs:
            r.font.size = Pt(12)
    tb.text_frame.word_wrap = True
for name in ("矩形 77", "矩形 80"):
    find(s, name).height = Inches(0.85)
for name in ("矩形 73", "矩形 74", "文本框 75", "椭圆 70"):
    el = find(s, name)._element
    el.getparent().remove(el)
notes(s, "Schedule. Until November: isolate the final check, implement AND/OR composition and run the agent on all "
         "specs of links A and B. December to January: add the OFDM receive and synchronization chains, the "
         "lowering agent, formal checks and the remaining baselines. February to March: main experiments and "
         "held-out re-validation. Then thesis writing and paper submission, with the defense expected in May 2027.")

# ---------------------------------------------------------------- drop old "Expected Outcomes" (slide 16)
sld_ids = prs.slides._sldIdLst
victim = sld_ids[15]
prs.part.drop_rel(victim.get(qn("r:id")))
sld_ids.remove(victim)

prs.save(OUT)
print("saved", OUT)
